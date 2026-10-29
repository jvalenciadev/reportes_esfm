import openpyxl
import json
import os

DEPT_MAP = {
    'chuquisaca': 1,
    'la paz': 2,
    'cochabamba': 3,
    'oruro': 4,
    'potosi': 5,
    'potosí': 5,
    'tarija': 6,
    'santa cruz': 7,
    'beni': 8,
    'pando': 9,
}

DEPT_NAMES = {
    1: 'Chuquisaca',
    2: 'La Paz',
    3: 'Cochabamba',
    4: 'Oruro',
    5: 'Potosí',
    6: 'Tarija',
    7: 'Santa Cruz',
    8: 'Beni',
    9: 'Pando',
}

def parse_num(v):
    if v is None:
        return 0
    if isinstance(v, (int, float)):
        return int(v)
    s = str(v).strip()
    if s in ('', '-', '–'):
        return 0
    try:
        return int(float(s.replace(',', '.')))
    except Exception:
        return 0

def clean_str(s):
    if s is None:
        return ''
    val = str(s).strip()
    val = val.replace('\ufffd', 'Í')
    return val

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
f1_path = os.path.join(BASE_DIR, 'public', 'datos', '-SEIE DISCAPACIDAD 2011-2025.xlsx')
f2_path = os.path.join(BASE_DIR, 'public', 'datos', '-SIE DISCAPACIDAD 2026.xlsx')

print(f'Leyendo {f1_path}...')
wb1 = openpyxl.load_workbook(f1_path, data_only=True)

# 1. FISCAL-PRIVADA
ws1 = wb1['FISCAL-PRIVADA']
rows_dep = []
for r in range(2, ws1.max_row + 1):
    dep = ws1.cell(r, 1).value
    est = ws1.cell(r, 2).value
    anio = ws1.cell(r, 3).value
    dept_name = ws1.cell(r, 4).value
    if dep is not None and anio is not None and dept_name is not None:
        d_clean = clean_str(dept_name)
        d_id = DEPT_MAP.get(d_clean.lower())
        if not d_id:
            continue
        rows_dep.append({
            'id': len(rows_dep) + 1,
            'departamento_id': d_id,
            'dependencia': clean_str(dep),
            'cantidad': parse_num(est),
            'anio': int(anio),
            'departamentos': {'nombre': DEPT_NAMES[d_id]}
        })

print(f'  1. seie_discapacidad_dependencia: {len(rows_dep)} filas')

# 2. RURAL-URBANA
ws2 = wb1['RURAL-URBANA']
rows_area = []
for r in range(2, ws2.max_row + 1):
    area = ws2.cell(r, 1).value
    est = ws2.cell(r, 2).value
    anio = ws2.cell(r, 3).value
    dept_name = ws2.cell(r, 4).value
    if area is not None and anio is not None and dept_name is not None:
        d_clean = clean_str(dept_name)
        d_id = DEPT_MAP.get(d_clean.lower())
        if not d_id:
            continue
        rows_area.append({
            'id': len(rows_area) + 1,
            'departamento_id': d_id,
            'area': clean_str(area),
            'cantidad': parse_num(est),
            'anio': int(anio),
            'departamentos': {'nombre': DEPT_NAMES[d_id]}
        })

print(f'  2. seie_discapacidad_area: {len(rows_area)} filas')

# 3. SIE DISCAPACIDAD 2026
print(f'Leyendo {f2_path}...')
wb2 = openpyxl.load_workbook(f2_path, data_only=True)
ws3 = wb2['Hoja1']
rows_sie2026 = []
for r in range(4, ws3.max_row + 1):
    dept_name = ws3.cell(r, 1).value
    dep = ws3.cell(r, 2).value
    est = ws3.cell(r, 3).value
    tipo = ws3.cell(r, 4).value
    if dept_name is not None and dep is not None and tipo is not None:
        d_clean = clean_str(dept_name)
        d_id = DEPT_MAP.get(d_clean.lower())
        if not d_id:
            continue
        rows_sie2026.append({
            'id': len(rows_sie2026) + 1,
            'departamento_id': d_id,
            'dependencia': clean_str(dep),
            'tipo_discapacidad': clean_str(tipo),
            'cantidad': parse_num(est),
            'anio': 2026,
            'departamentos': {'nombre': DEPT_NAMES[d_id]}
        })

print(f'  3. sie_discapacidad_2026: {len(rows_sie2026)} filas')

# Actualizar fallbackData.json
fallback_path = os.path.join(BASE_DIR, 'src', 'data', 'fallbackData.json')
print(f'Actualizando {fallback_path}...')
with open(fallback_path, 'r', encoding='utf-8') as f:
    fb = json.load(f)

fb['seie_discapacidad_dependencia'] = rows_dep
fb['seie_discapacidad_area'] = rows_area
fb['sie_discapacidad_2026'] = rows_sie2026

with open(fallback_path, 'w', encoding='utf-8') as f:
    json.dump(fb, f, ensure_ascii=False)
print('fallbackData.json actualizado correctamente.')

# Generar scripts/discapacidad_v3.sql
sql_path = os.path.join(BASE_DIR, 'scripts', 'discapacidad_v3.sql')
print(f'Generando {sql_path}...')
with open(sql_path, 'w', encoding='utf-8') as f:
    f.write('''-- MIGRACIÓN DE DATOS DE EDUCACIÓN ESPECIAL Y DISCAPACIDAD (3 TABLAS)
-- 1. seie_discapacidad_dependencia (2011 - 2025: Fiscal vs Privada)
-- 2. seie_discapacidad_area (2011 - 2025: Rural vs Urbana)
-- 3. sie_discapacidad_2026 (2026: Por tipo de discapacidad y dependencia)

-- ============================================================================
-- TABLA 1: seie_discapacidad_dependencia
-- ============================================================================
CREATE TABLE IF NOT EXISTS public.seie_discapacidad_dependencia (
    id BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    departamento_id INT NOT NULL REFERENCES public.departamentos(id) ON DELETE CASCADE,
    dependencia VARCHAR(50) NOT NULL,
    cantidad INT NOT NULL,
    anio INT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_seie_disc_dep_anio ON public.seie_discapacidad_dependencia(anio, departamento_id);
TRUNCATE TABLE public.seie_discapacidad_dependencia RESTART IDENTITY CASCADE;

''')

    # Batch inserts para tabla 1
    chunk_size = 100
    for i in range(0, len(rows_dep), chunk_size):
        chunk = rows_dep[i:i+chunk_size]
        vals = [f"({r['departamento_id']}, '{r['dependencia']}', {r['cantidad']}, {r['anio']})" for r in chunk]
        f.write(f"INSERT INTO public.seie_discapacidad_dependencia (departamento_id, dependencia, cantidad, anio) VALUES\n  " + ",\n  ".join(vals) + ";\n\n")

    f.write('''-- ============================================================================
-- TABLA 2: seie_discapacidad_area
-- ============================================================================
CREATE TABLE IF NOT EXISTS public.seie_discapacidad_area (
    id BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    departamento_id INT NOT NULL REFERENCES public.departamentos(id) ON DELETE CASCADE,
    area VARCHAR(50) NOT NULL,
    cantidad INT NOT NULL,
    anio INT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_seie_disc_area_anio ON public.seie_discapacidad_area(anio, departamento_id);
TRUNCATE TABLE public.seie_discapacidad_area RESTART IDENTITY CASCADE;

''')

    # Batch inserts para tabla 2
    for i in range(0, len(rows_area), chunk_size):
        chunk = rows_area[i:i+chunk_size]
        vals = [f"({r['departamento_id']}, '{r['area']}', {r['cantidad']}, {r['anio']})" for r in chunk]
        f.write(f"INSERT INTO public.seie_discapacidad_area (departamento_id, area, cantidad, anio) VALUES\n  " + ",\n  ".join(vals) + ";\n\n")

    f.write('''-- ============================================================================
-- TABLA 3: sie_discapacidad_2026
-- ============================================================================
CREATE TABLE IF NOT EXISTS public.sie_discapacidad_2026 (
    id BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    departamento_id INT NOT NULL REFERENCES public.departamentos(id) ON DELETE CASCADE,
    dependencia VARCHAR(50) NOT NULL,
    tipo_discapacidad VARCHAR(100) NOT NULL,
    cantidad INT NOT NULL,
    anio INT NOT NULL DEFAULT 2026
);
CREATE INDEX IF NOT EXISTS idx_sie_disc_2026_dept ON public.sie_discapacidad_2026(departamento_id);
TRUNCATE TABLE public.sie_discapacidad_2026 RESTART IDENTITY CASCADE;

''')

    # Batch inserts para tabla 3
    for i in range(0, len(rows_sie2026), chunk_size):
        chunk = rows_sie2026[i:i+chunk_size]
        vals = [f"({r['departamento_id']}, '{r['dependencia']}', '{r['tipo_discapacidad']}', {r['cantidad']}, {r['anio']})" for r in chunk]
        f.write(f"INSERT INTO public.sie_discapacidad_2026 (departamento_id, dependencia, tipo_discapacidad, cantidad, anio) VALUES\n  " + ",\n  ".join(vals) + ";\n\n")

    f.write('''-- Habilitar RLS en las 3 tablas
DO $$
DECLARE
    tbl TEXT;
BEGIN
    FOR tbl IN SELECT unnest(ARRAY[
        'seie_discapacidad_dependencia',
        'seie_discapacidad_area',
        'sie_discapacidad_2026'
    ])
    LOOP
        EXECUTE format('ALTER TABLE public.%I ENABLE ROW LEVEL SECURITY;', tbl);
        EXECUTE format('DROP POLICY IF EXISTS "allow_anon_all" ON public.%I;', tbl);
        EXECUTE format('CREATE POLICY "allow_anon_all" ON public.%I FOR ALL TO anon, authenticated USING (true) WITH CHECK (true);', tbl);
    END LOOP;
END $$;
''')

print(f'discapacidad_v3.sql generado exitosamente en {sql_path}')

