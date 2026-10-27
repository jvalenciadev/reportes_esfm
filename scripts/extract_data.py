# -*- coding: utf-8 -*-
"""
Extrae datos de los 7 archivos Excel y genera un JSON limpio (UTF-8).
Corrige los caracteres corruptos (U+FFFD - Unicode Replacement Character)
que aparecen en lugar de tildes y eñes en los archivos xlsx.
"""
import openpyxl
import json
import re
import os
import sys
import argparse

parser = argparse.ArgumentParser()
parser.add_argument('--out', default=None, help='Output file path')
args, _ = parser.parse_known_args()

FOLDER = r"c:\Users\PROFE-JP\Desktop\PROYECTO\esfm_admision_reporte\public\datos"
R = '\ufffd'  # Unicode Replacement Character - aparece en lugar de tildes

# Pares (texto_corrupto, texto_correcto)
FIXES = [
    (f'PL{R}STICAS', 'PLÁSTICAS'),
    (f'CARPINTER{R}A', 'CARPINTERÍA'),
    (f'BIOLOG{R}A', 'BIOLOGÍA'),
    (f'GEOGRAF{R}A', 'GEOGRAFÍA'),
    (f'F{R}SICA', 'FÍSICA'),
    (f'QU{R}MICA', 'QUÍMICA'),
    (f'MENCI{R}N', 'MENCIÓN'),
    (f'COMUNICACI{R}N', 'COMUNICACIÓN'),
    (f'INGL{R}S', 'INGLÉS'),
    (f'FILOSOF{R}AS', 'FILOSOFÍAS'),
    (f'FILOSOF{R}A', 'FILOSOFÍA'),
    (f'SICOLOG{R}A', 'SICOLOGÍA'),
    (f'TECNOL{R}GICA', 'TECNOLÓGICA'),
    (f'TECNOL{R}GICO', 'TECNOLÓGICO'),
    (f'HUMAN{R}STICO', 'HUMANÍSTICO'),
    (f'EDUCACI{R}N', 'EDUCACIÓN'),
    (f'EDUCACION', 'EDUCACIÓN'),
    (f'T{R}CNICA', 'TÉCNICA'),
    (f'ELECTR{R}NICA', 'ELECTRÓNICA'),
    (f'MEC{R}NICA', 'MECÁNICA'),
    (f'MATEM{R}TICA', 'MATEMÁTICA'),
    (f'SOBERAN{R}A', 'SOBERANÍA'),
    (f'TRANSFORMACI{R}N', 'TRANSFORMACIÓN'),
    (f'GASTRONOM{R}A', 'GASTRONOMÍA'),
    (f'GASTRONOMIA', 'GASTRONOMÍA'),
    (f'CONFECCI{R}N', 'CONFECCIÓN'),
    (f'P{R}BLICO', 'PÚBLICO'),
    (f'P{R}blica', 'Pública'),
    (f'P{R}blico', 'Público'),
    (f'POTOS{R}', 'POTOSÍ'),
    (f'Potos{R}', 'Potosí'),
    (f'POBLACI{R}N', 'POBLACIÓN'),
    (f'JOS{R} DAVID', 'JOSÉ DAVID'),
    (f'JOS{R}', 'JOSÉ'),
    (f'ANDR{R}S', 'ANDRÉS'),
    (f'MULTI{R}TNICA', 'MULTIÉTNICA'),
    (f'IND{R}GENA', 'INDÍGENA'),
    (f'CONCEPCI{R}N', 'CONCEPCIÓN'),
    (f'PLURI{R}TNICA', 'PLURIÉTNICA'),
    (f'CH{R}VEZ', 'CHÁVEZ'),
    (f'SIM{R}N', 'SIMÓN'),
    (f'{R}NGEL', 'ÁNGEL'),
    (f'CIENCIAS  SOCIALES', 'CIENCIAS SOCIALES'),
    # Dependencia en SIE
    (f'P{R}blico  ', 'Público'),
    (f'Privado ', 'Privado'),
]

def clean(v):
    if v is None:
        return ''
    s = str(v).strip()
    for bad, good in FIXES:
        s = s.replace(bad, good)
    return re.sub(r'\s+', ' ', s).strip()

def parse_int(v, default=0):
    if v is None:
        return default
    s = str(v).strip()
    if s in ('', '-', 'None', 'null', 'N/A'):
        return default
    try:
        return int(float(s.replace(',', '')))
    except Exception:
        return default

DEP_MAP = {
    'CHUQUISACA': 1, 'LA PAZ': 2, 'COCHABAMBA': 3,
    'ORURO': 4, 'POTOSI': 5, 'POTOSÍ': 5,
    'TARIJA': 6, 'SANTA CRUZ': 7, 'BENI': 8, 'PANDO': 9
}

def load(filename):
    wb = openpyxl.load_workbook(os.path.join(FOLDER, filename), data_only=True)
    return wb.active

result = {}

# -----------------------------------
# 1. Catálogo ESPECIALIDADES
# -----------------------------------
sh = load('-INSCRITOS ESFM 2016 - 2026.xlsx')
esps = sorted(set(
    clean(sh.cell(r, 4).value)
    for r in range(2, sh.max_row + 1)
    if sh.cell(r, 4).value
))
esp_map = {n: i + 1 for i, n in enumerate(esps)}
result['especialidades'] = [{'id': v, 'nombre': k} for k, v in esp_map.items()]

# -----------------------------------
# 2. Catálogo ESFM
# -----------------------------------
sh = load('-INSCRITOS ESFM 2016 - 2026.xlsx')
esfms = sorted(set(
    clean(sh.cell(r, 3).value)
    for r in range(2, sh.max_row + 1)
    if sh.cell(r, 3).value
))
esfm_map = {n: i + 1 for i, n in enumerate(esfms)}
result['esfm'] = [{'id': v, 'nombre': k} for k, v in esfm_map.items()]

# -----------------------------------
# 3. DISCAPACIDAD SEIE
# -----------------------------------
sh = load('-DISCAPACIDAD SEIE 2011-2025.xlsx')
rows = []
for r in range(2, sh.max_row + 1):
    area = clean(sh.cell(r, 1).value)
    cant = parse_int(sh.cell(r, 2).value)
    anio = parse_int(sh.cell(r, 3).value)
    if area and anio:
        rows.append({'area': area, 'cantidad_est': cant, 'anio': anio})
result['discapacidad_seie'] = rows

# -----------------------------------
# 4. INE MATRÍCULA
# -----------------------------------
sh = load('-INE MATRICULA 2011-2025.xlsx')
rows = []
for r in range(2, sh.max_row + 1):
    dep_id = DEP_MAP.get(clean(sh.cell(r, 1).value).upper())
    tipo = clean(sh.cell(r, 2).value)
    nivel = clean(sh.cell(r, 3).value)
    cant = parse_int(sh.cell(r, 4).value)
    anio = parse_int(sh.cell(r, 5).value)
    if dep_id and tipo and anio:
        rows.append({'departamento_id': dep_id, 'tipo': tipo, 'nivel': nivel, 'cantidad': cant, 'anio': anio})
result['ine_matricula'] = rows

# -----------------------------------
# 5. INE POBLACIÓN
# -----------------------------------
sh = load('-INE POBLACION 2000- 2050.xlsx')
rows = []
for r in range(2, sh.max_row + 1):
    anio = parse_int(sh.cell(r, 1).value)
    pob = parse_int(sh.cell(r, 2).value)
    if anio:
        rows.append({'anio': anio, 'poblacion': pob})
result['ine_poblacion'] = rows

# -----------------------------------
# 6. INE PROYECCIÓN POR EDAD
# -----------------------------------
sh = load('-INE PROYECCION POR EDAD 2024-2034.xlsx')
rows = []
for r in range(2, sh.max_row + 1):
    edad = parse_int(sh.cell(r, 1).value)
    cant = parse_int(sh.cell(r, 2).value)
    anio = parse_int(sh.cell(r, 3).value)
    if anio:
        rows.append({'edad': edad, 'cantidad': cant, 'anio': anio})
result['ine_proyeccion_edad'] = rows

# -----------------------------------
# 7. INSCRITOS ESFM (con esfm_id)
# -----------------------------------
sh = load('-INSCRITOS ESFM 2016 - 2026.xlsx')
rows = []
skipped = 0
for r in range(2, sh.max_row + 1):
    gestion = parse_int(sh.cell(r, 1).value)
    dep_id = DEP_MAP.get(clean(sh.cell(r, 2).value).upper())
    esfm_name = clean(sh.cell(r, 3).value)
    esfm_id = esfm_map.get(esfm_name)
    esp_name = clean(sh.cell(r, 4).value)
    esp_id = esp_map.get(esp_name)
    modalidad = clean(sh.cell(r, 5).value)
    tot = parse_int(sh.cell(r, 6).value)
    if gestion and dep_id and esfm_id and esp_id:
        rows.append({
            'gestion': gestion,
            'departamento_id': dep_id,
            'esfm_id': esfm_id,
            'especialidad_id': esp_id,
            'modalidad': modalidad,
            'total_inscritos': tot
        })
    else:
        skipped += 1
result['inscritos_esfm'] = rows
result['_inscritos_skipped'] = skipped

# -----------------------------------
# 8. MATRÍCULA RURAL URBANA
# -----------------------------------
sh = load('-MATRI ESC RURAL URBANA 2000-2025.xlsx')
rows = []
for r in range(2, sh.max_row + 1):
    area = clean(sh.cell(r, 1).value)
    cant = parse_int(sh.cell(r, 2).value)
    anio = parse_int(sh.cell(r, 3).value)
    if area and anio:
        rows.append({'area': area, 'cantidad': cant, 'anio': anio})
result['matricula_rural_urbana'] = rows

# -----------------------------------
# 9. SIE TOTAL ESTUDIANTES
# -----------------------------------
sh = load('-SIE TOTAL EST INC-PRI-SEC 2026.xlsx')
rows = []
for r in range(2, sh.max_row + 1):
    dep_id = DEP_MAP.get(clean(sh.cell(r, 1).value).upper())
    depen = clean(sh.cell(r, 2).value)
    cant = parse_int(sh.cell(r, 3).value)
    nivel = parse_int(sh.cell(r, 4).value)
    nivel_txt = clean(sh.cell(r, 5).value)
    if dep_id:
        rows.append({'departamento_id': dep_id, 'dependencia': depen, 'cantidad': cant, 'nivel': nivel, 'nivel_texto': nivel_txt})
result['sie_total_estudiantes'] = rows

# Estadísticas
stats = {k: len(v) if isinstance(v, list) else v for k, v in result.items()}
sys.stderr.write(f"Estadísticas: {json.dumps(stats, ensure_ascii=False)}\n")

output_json = json.dumps(result, ensure_ascii=False, indent=None)

if args.out:
    with open(args.out, 'w', encoding='utf-8') as f:
        f.write(output_json)
    sys.stderr.write(f"JSON guardado en: {args.out}\n")
else:
    sys.stdout.buffer.write(output_json.encode('utf-8'))
