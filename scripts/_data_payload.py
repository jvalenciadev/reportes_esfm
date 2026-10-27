
import openpyxl, json, re, os
folder = r'c:\\Users\\PROFE-JP\\Desktop\\PROYECTO\\esfm_admision_reporte\\public\\datos'

fixes = {
  'PLSTICAS': 'PLÁSTICAS', 'CARPINTERA': 'CARPINTERÍA',
  'BIOLOGA': 'BIOLOGÍA', 'GEOGRAFA': 'GEOGRAFÍA',
  'FSICA': 'FÍSICA', 'QUMICA': 'QUÍMICA',
  'MENCIN': 'MENCIÓN', 'COMUNICACIN': 'COMUNICACIÓN',
  'INGLS': 'INGLÉS', 'FILOSOFA': 'FILOSOFÍA',
  'SICOLOGA': 'SICOLOGÍA', 'FILOSOFS': 'FILOSOFÍAS',
  'TECNOLGICA': 'TECNOLÓGICA', 'EDUCACIN': 'EDUCACIÓN',
  'EDUCACION': 'EDUCACIÓN', 'TCNICA': 'TÉCNICA',
  'ELECTRNICA': 'ELECTRÓNICA', 'MECNICA': 'MECÁNICA',
  'MATEMTICA': 'MATEMÁTICA', 'SOBERANA': 'SOBERANÍA',
  'TRANSFORMACIN': 'TRANSFORMACIÓN', 'GASTRONOMA': 'GASTRONOMÍA',
  'GASTRONOMIA': 'GASTRONOMÍA', 'CONFECCIN': 'CONFECCIÓN',
  'PBLICO': 'PÚBLICO', 'POTOS': 'POTOSÍ',
  'POBLACIN': 'POBLACIÓN', 'JOS': 'JOSÉ',
  'ANDRS': 'ANDRÉS', 'MULTITNICA': 'MULTIÉTNICA',
  'INDGENA': 'INDÍGENA', 'CONCEPCIN': 'CONCEPCIÓN',
  'PLURITNICA': 'PLURIÉTNICA', 'CHVEZ': 'CHÁVEZ',
  'SIMN': 'SIMÓN', 'NGEL': 'ÁNGEL',
}

def c(v):
    if v is None: return ''
    s = str(v).strip()
    for bad, good in fixes.items():
        s = s.replace(bad, good)
    return re.sub(r'\\s+', ' ', s).strip()

dep_map = {'CHUQUISACA':1,'LA PAZ':2,'COCHABAMBA':3,'ORURO':4,'POTOSI':5,'POTOSÍ':5,'TARIJA':6,'SANTA CRUZ':7,'BENI':8,'PANDO':9}
pi = lambda v, d=0: d if v is None or str(v).strip() in ('','-','None') else int(float(str(v).replace(',','')))
pb = lambda v, d=0: d if v is None or str(v).strip() in ('','-','None') else int(float(str(v).replace(',','')))

result = {}

# Especialidades
wb = openpyxl.load_workbook(os.path.join(folder, '-INSCRITOS ESFM 2016 - 2026.xlsx'), data_only=True)
sh = wb.active
esps = sorted(set(c(sh.cell(r,4).value) for r in range(2, sh.max_row+1) if sh.cell(r,4).value))
esp_map = {n:i+1 for i,n in enumerate(esps)}
result['especialidades'] = [{'id':i,'nombre':n} for n,i in esp_map.items()]

# Discapacidad SEIE
wb = openpyxl.load_workbook(os.path.join(folder, '-DISCAPACIDAD SEIE 2011-2025.xlsx'), data_only=True)
sh = wb.active
rows = []
for r in range(2, sh.max_row+1):
    area, cant, anio = c(sh.cell(r,1).value), pi(sh.cell(r,2).value), pi(sh.cell(r,3).value)
    if area and anio: rows.append({'area':area,'cantidad_est':cant,'anio':anio})
result['discapacidad_seie'] = rows

# INE Matricula
wb = openpyxl.load_workbook(os.path.join(folder, '-INE MATRICULA 2011-2025.xlsx'), data_only=True)
sh = wb.active
rows = []
for r in range(2, sh.max_row+1):
    dep_id = dep_map.get(c(sh.cell(r,1).value).upper())
    tipo, nivel, cant, anio = c(sh.cell(r,2).value), c(sh.cell(r,3).value), pi(sh.cell(r,4).value), pi(sh.cell(r,5).value)
    if dep_id and tipo and anio: rows.append({'departamento_id':dep_id,'tipo':tipo,'nivel':nivel,'cantidad':cant,'anio':anio})
result['ine_matricula'] = rows

# INE Poblacion
wb = openpyxl.load_workbook(os.path.join(folder, '-INE POBLACION 2000- 2050.xlsx'), data_only=True)
sh = wb.active
rows = []
for r in range(2, sh.max_row+1):
    anio, pob = pi(sh.cell(r,1).value), pb(sh.cell(r,2).value)
    if anio: rows.append({'anio':anio,'poblacion':pob})
result['ine_poblacion'] = rows

# INE Proyeccion por edad
wb = openpyxl.load_workbook(os.path.join(folder, '-INE PROYECCION POR EDAD 2024-2034.xlsx'), data_only=True)
sh = wb.active
rows = []
for r in range(2, sh.max_row+1):
    edad, cant, anio = pi(sh.cell(r,1).value), pi(sh.cell(r,2).value), pi(sh.cell(r,3).value)
    if anio: rows.append({'edad':edad,'cantidad':cant,'anio':anio})
result['ine_proyeccion_edad'] = rows

# Inscritos ESFM
wb = openpyxl.load_workbook(os.path.join(folder, '-INSCRITOS ESFM 2016 - 2026.xlsx'), data_only=True)
sh = wb.active
rows = []
for r in range(2, sh.max_row+1):
    gest = pi(sh.cell(r,1).value)
    dep_id = dep_map.get(c(sh.cell(r,2).value).upper())
    esfm = c(sh.cell(r,3).value)
    esp_id = esp_map.get(c(sh.cell(r,4).value))
    mod, tot = c(sh.cell(r,5).value), pi(sh.cell(r,6).value)
    if gest and dep_id and esp_id: rows.append({'gestion':gest,'departamento_id':dep_id,'esfm':esfm,'especialidad_id':esp_id,'modalidad':mod,'total_inscritos':tot})
result['inscritos_esfm'] = rows

# Matricula Rural Urbana
wb = openpyxl.load_workbook(os.path.join(folder, '-MATRI ESC RURAL URBANA 2000-2025.xlsx'), data_only=True)
sh = wb.active
rows = []
for r in range(2, sh.max_row+1):
    area, cant, anio = c(sh.cell(r,1).value), pi(sh.cell(r,2).value), pi(sh.cell(r,3).value)
    if area and anio: rows.append({'area':area,'cantidad':cant,'anio':anio})
result['matricula_rural_urbana'] = rows

# SIE Total Estudiantes
wb = openpyxl.load_workbook(os.path.join(folder, '-SIE TOTAL EST INC-PRI-SEC 2026.xlsx'), data_only=True)
sh = wb.active
rows = []
for r in range(2, sh.max_row+1):
    dep_id = dep_map.get(c(sh.cell(r,1).value).upper())
    depen, cant, nivel, nivel_txt = c(sh.cell(r,2).value), pi(sh.cell(r,3).value), pi(sh.cell(r,4).value), c(sh.cell(r,5).value)
    if dep_id: rows.append({'departamento_id':dep_id,'dependencia':depen,'cantidad':cant,'nivel':nivel,'nivel_texto':nivel_txt})
result['sie_total_estudiantes'] = rows

print(json.dumps(result, ensure_ascii=False))
